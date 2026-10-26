"""Apply Siddharth's decisions from the Q&A review spreadsheet to the website chatbot's bank (sam_web_qa).

  python prototype/import_qa_review.py --dry-run          # read the sheet and the bank (GETs only), print the plan, write nothing
  python prototype/import_qa_review.py                    # apply it
  python prototype/import_qa_review.py FILE.xlsx --by siddharth
  python prototype/import_qa_review.py --self-check       # offline asserts, no network

Default file: ~/Downloads/SAM-Website-Bot-QA-Review-2026-10-26.xlsx, sheet "Q&A review". Columns are found by their
header, in any order: ID, Question, Answer the bot shows, Visual, Needs your sign-off because, Evidence, Source pages,
Decision, Your edit / comment.

  Decision  What is written (the same shape as the admin Website tab, web/lib/web-dashboard.ts decisionPatch)
  Approve   status 'approved', approved_by, approved_at (now), review_by (12 months on, India date)
  Edit      answer := "Your edit / comment", then approved as above
  Reject    status 'retired' (never shown; it can be approved again in the tab)
  (blank)   left alone

Safety, in order:
  - Every row is checked before the first write; one bad row (an unknown decision, an edit over 80 words or 600
    characters, no source page, an empty edit) stops the whole run. Nothing is half-applied. An id that is not in the
    bank (lp-27: the loader refused it) is skipped with a loud note.
  - Approve refuses a pair whose answer in the database is no longer the one in the sheet: it changed after the sheet
    was made, so what was read is not what would go public. Review it again (or use Edit).
  - Each write is a PATCH filtered on the pair's current status, so a decision made in the tab meanwhile is not
    overwritten. A pair already in the target state is skipped; re-running is harmless.
Credentials: web/.env.local (SUPABASE_URL / SUPABASE_SERVICE_KEY), as load_web_knowledge.py. The approver defaults to
the first admin login in SAM_USERS (the name the tab records); passwords and keys are never printed.
"""
from __future__ import annotations
import argparse, calendar, json, re, sys, urllib.error, urllib.parse, urllib.request
from datetime import datetime, timedelta, timezone
from pathlib import Path

HERE = Path(__file__).parent
REPO = HERE.parent
ENV = REPO / "web" / ".env.local"
DEFAULT_XLSX = Path.home() / "Downloads" / "SAM-Website-Bot-QA-Review-2026-10-26.xlsx"
SHEET = "Q&A review"
MAX_WORDS, MAX_CHARS, REVIEW_MONTHS = 80, 600, 12
IST = timezone(timedelta(hours=5, minutes=30))
ID_RE = re.compile(r"^[a-z0-9][a-z0-9-]{0,79}$")
DECISIONS = {"approve": "approve", "approved": "approve", "yes": "approve", "edit": "edit", "edited": "edit",
             "reject": "retire", "rejected": "retire", "retire": "retire", "no": "retire"}
COLS = {"id": "id", "question": "question", "answer the bot shows": "answer", "decision": "decision", "your edit": "edit"}


def ws(s: str | None) -> str:
    return re.sub(r"\s+", " ", (s or "").strip())


def answer_problem(text: str) -> str | None:
    """The table's own check (01: <= 600 chars, <= 80 words), in a sentence."""
    t = text.strip()
    if not t:
        return "the edit is empty"
    if len(t) > MAX_CHARS:
        return f"the edit is {len(t)} characters; the most is {MAX_CHARS}"
    if len(t.split()) > MAX_WORDS:
        return f"the edit is {len(t.split())} words; the most is {MAX_WORDS}"
    return None


def review_by(now: datetime) -> str:
    """India date + 12 months, the day clamped to the month's end (29 Feb -> 28 Feb). Same rule as reviewBy() in TS."""
    d = now.astimezone(IST).date()
    m0 = d.month - 1 + REVIEW_MONTHS
    y, m = d.year + m0 // 12, m0 % 12 + 1
    return f"{y:04d}-{m:02d}-{min(d.day, calendar.monthrange(y, m)[1]):02d}"


def read_sheet(path: Path) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sh = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    rows = list(sh.iter_rows(values_only=True))
    wb.close()                                   # read-only mode holds the file open (Windows locks it)
    if not rows:
        return []
    head = [ws(str(h or "")).lower() for h in rows[0]]
    idx = {}
    for i, h in enumerate(head):
        for prefix, key in COLS.items():
            if h.startswith(prefix) and key not in idx:
                idx[key] = i
    missing = [k for k in ("id", "decision", "edit", "answer") if k not in idx]
    if missing:
        sys.exit(f"{path.name}: no column for {', '.join(missing)} (headers: {', '.join(head)})")
    out = []
    for n, r in enumerate(rows[1:], start=2):
        get = lambda k: (str(r[idx[k]]) if idx.get(k) is not None and idx[k] < len(r) and r[idx[k]] is not None else "")
        if not ws(get("id")):
            continue
        out.append({"row": n, "id": ws(get("id")), "question": ws(get("question")), "answer": get("answer"),
                    "decision": ws(get("decision")), "edit": get("edit")})
    return out


def plan(sheet: list[dict], db: dict[str, dict], by: str, now: datetime) -> tuple[list[dict], list[str], list[str]]:
    """-> (writes, problems, notes). Writes are {id, from, patch, what}. Any problem means nothing is written."""
    writes, problems, notes = [], [], []
    seen = set()
    stamp = {"status": "approved", "approved_by": by, "approved_at": now.astimezone(timezone.utc).isoformat(), "review_by": review_by(now)}
    for s in sheet:
        where = f"row {s['row']} ({s['id']})"
        if not s["decision"]:
            continue
        d = DECISIONS.get(s["decision"].lower())
        if not d:
            problems.append(f"{where}: unknown decision '{s['decision']}' (Approve, Edit or Reject)")
            continue
        if not ID_RE.match(s["id"]) or s["id"] not in db:   # e.g. lp-27: in the sheet, refused by the loader (blocked claim)
            notes.append(f"SKIPPED {where}: not in sam_web_qa (never loaded), so '{s['decision']}' cannot apply")
            continue
        if s["id"] in seen:
            problems.append(f"{where}: this id has a decision on another row too")
            continue
        seen.add(s["id"])
        cur = db[s["id"]]
        if d == "retire":
            if cur["status"] == "retired":
                notes.append(f"{s['id']}: already retired, skipped")
            else:
                writes.append({"id": s["id"], "from": cur["status"], "patch": {"status": "retired"}, "what": "retire"})
            continue
        if not cur.get("sources"):
            problems.append(f"{where}: the pair has no source page, so it cannot be approved")
            continue
        if d == "approve":
            if ws(s["edit"]):
                notes.append(f"{s['id']}: comment kept in the sheet only (Approve ignores it): {ws(s['edit'])[:80]}")
            if ws(cur["answer"]) != ws(s["answer"]):
                problems.append(f"{where}: the answer changed after the sheet was made; read the current one in the Website tab "
                                f"(or put the wording you want in 'Your edit' and choose Edit)")
                continue
            if cur["status"] == "approved":
                notes.append(f"{s['id']}: already approved, skipped")
            else:
                writes.append({"id": s["id"], "from": cur["status"], "patch": dict(stamp), "what": "approve"})
            continue
        p = answer_problem(s["edit"])
        if p:
            problems.append(f"{where}: {p}")
            continue
        new = ws(s["edit"])
        if new == ws(cur["answer"]) and cur["status"] == "approved":
            notes.append(f"{s['id']}: edit equals the approved answer, skipped")
            continue
        writes.append({"id": s["id"], "from": cur["status"], "patch": {"answer": new, **stamp}, "what": f"edit ({len(new.split())} words) + approve"})
    return writes, problems, notes


def env(path: Path = ENV) -> dict[str, str]:
    """web/.env.local, then the process environment for anything it lacks."""
    import os
    out = {}
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            if "=" in line and not line.lstrip().startswith("#"):
                k, v = line.split("=", 1)
                out[k.strip()] = v.strip().strip('"')
    for k in ("SUPABASE_URL", "SUPABASE_SERVICE_KEY", "SAM_USERS"):
        if not out.get(k) and os.environ.get(k):
            out[k] = os.environ[k]
    return out


def default_by(e: dict[str, str]) -> str | None:
    for part in (e.get("SAM_USERS") or "").split(","):
        bits = [b.strip() for b in part.split(":")]
        if len(bits) >= 3 and bits[2] == "admin" and bits[0]:
            return bits[0].lower()
    return None


class Db:
    def __init__(self, url: str, key: str):
        self.url, self.key = url.rstrip("/"), key

    def call(self, method: str, path: str, body=None, prefer: str | None = None):
        h = {"apikey": self.key, "Authorization": f"Bearer {self.key}", "Content-Type": "application/json"}
        if prefer:
            h["Prefer"] = prefer
        req = urllib.request.Request(f"{self.url}/rest/v1/{path}", method=method, headers=h,
                                     data=None if body is None else json.dumps(body).encode())
        try:
            with urllib.request.urlopen(req, timeout=20) as r:
                t = r.read().decode()
                return json.loads(t) if t else None
        except urllib.error.HTTPError as e:   # status + PostgREST's message only; never the key
            raise SystemExit(f"{method} {path.split('?')[0]} returned {e.code}: {e.read().decode()[:200]}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--by", help="approver login recorded in approved_by (default: the first admin in SAM_USERS)")
    ap.add_argument("--env", type=Path, default=ENV, help="the .env file with SUPABASE_URL / SUPABASE_SERVICE_KEY (default web/.env.local)")
    ap.add_argument("--self-check", action="store_true")
    a = ap.parse_args()
    if a.self_check:
        return self_check()
    e = env(a.env)
    by = (a.by or default_by(e) or "").strip().lower()
    if not by:
        sys.exit("no approver: pass --by <your admin login>")
    if not a.xlsx.exists():
        sys.exit(f"no file {a.xlsx}")
    url, key = e.get("SUPABASE_URL"), e.get("SUPABASE_SERVICE_KEY")
    if not url or not key:
        sys.exit(f"SUPABASE_URL / SUPABASE_SERVICE_KEY missing from {a.env}")
    sheet = read_sheet(a.xlsx)
    decided = [s for s in sheet if s["decision"]]
    print(f"{a.xlsx.name}: {len(sheet)} pairs, {len(decided)} with a decision, {len(sheet) - len(decided)} undecided (left alone). Approver: {by}")
    db = Db(url, key)
    ids = sorted({s["id"] for s in decided if ID_RE.match(s["id"])})
    rows = db.call("GET", "sam_web_qa?select=id,status,answer,sources&id=in." + urllib.parse.quote(f"({','.join(ids)})")) if ids else []
    writes, problems, notes = plan(sheet, {r["id"]: r for r in rows}, by, datetime.now(timezone.utc))
    for w in writes:
        print(f"  {w['id']:<14} {w['from']:>8} -> {w['what']}")
    for n in notes:
        print(f"  note: {n}")
    if problems:
        print(f"\n{len(problems)} problem(s); nothing {'would be' if a.dry_run else 'was'} written:")
        for p in problems:
            print(f"  - {p}")
        sys.exit(1)
    if a.dry_run:
        print(f"\ndry run: {len(writes)} write(s) planned, none made. Run again without --dry-run to apply.")
        return
    done = 0
    for w in writes:
        got = db.call("PATCH", f"sam_web_qa?id=eq.{urllib.parse.quote(w['id'])}&status=eq.{w['from']}", w["patch"], "return=representation")
        if len(got or []) == 1:
            done += 1
        else:
            print(f"  {w['id']}: changed by someone else meanwhile (status no longer '{w['from']}'), skipped")
    print(f"applied {done} of {len(writes)}. Clicks and chips show approved pairs within 5 minutes; "
          f"rebuild the intent index so typed questions find them: python prototype/build_intent_index.py")


def self_check() -> None:
    now = datetime(2026, 10, 26, 20, 0, tzinfo=timezone.utc)                       # 27 Oct 01:30 IST
    assert review_by(now) == "2027-10-27", review_by(now)
    assert review_by(datetime(2028, 2, 29, 6, 0, tzinfo=timezone.utc)) == "2029-02-28"
    assert answer_problem("") and answer_problem("w " * 81) and answer_problem("x" * 601) and answer_problem("Fine.") is None
    db = {"lp-01": {"id": "lp-01", "status": "draft", "answer": "Yes,  it is.", "sources": ["https://a"]},
          "lp-02": {"id": "lp-02", "status": "draft", "answer": "Old.", "sources": ["https://a"]},
          "lp-03": {"id": "lp-03", "status": "approved", "answer": "A.", "sources": ["https://a"]},
          "lp-04": {"id": "lp-04", "status": "draft", "answer": "B.", "sources": []}}
    row = lambda i, d, ans="", edit="", n=2: {"row": n, "id": i, "question": "", "answer": ans, "decision": d, "edit": edit}
    w, p, n = plan([row("lp-01", "Approve", "Yes, it is."), row("lp-02", "edit", "", "New answer here."), row("lp-03", "Reject"),
                    row("lp-05", ""), row("lp-03", "")], db, "siddharth", now)
    assert not p, p
    assert [x["what"] for x in w] == ["approve", "edit (3 words) + approve", "retire"], w
    assert w[0]["patch"] == {"status": "approved", "approved_by": "siddharth", "approved_at": now.isoformat(), "review_by": "2027-10-27"}
    assert w[1]["patch"]["answer"] == "New answer here." and w[1]["patch"]["status"] == "approved" and w[2]["patch"] == {"status": "retired"}
    for bad, want in [(row("lp-01", "Approve", "Something else."), "answer changed"), (row("lp-02", "Edit", "", "w " * 90), "90 words"),
                      (row("lp-02", "Edit", "", " "), "empty"), (row("lp-01", "Maybe"), "unknown decision"),
                      (row("lp-04", "Approve", "B."), "no source")]:
        _, p, _ = plan([bad], db, "siddharth", now)
        assert p and want in p[0], (want, p)
    w, p, n = plan([row("lp-27", "Approve"), row("lp-01", "Approve", "Yes, it is.")], db, "s", now)
    assert not p and len(w) == 1 and n[0].startswith("SKIPPED row 2 (lp-27)"), (p, n)
    _, p, _ = plan([row("lp-02", "Edit", "", "One."), row("lp-02", "Approve", "Old.", n=3)], db, "s", now)
    assert p and "another row" in p[0], p
    w, p, n = plan([row("lp-03", "Approve", "A.")], db, "s", now)
    assert not w and not p and "already approved" in n[0]
    # the sheet reader, on a real workbook in a temp dir
    import openpyxl, tempfile
    with tempfile.TemporaryDirectory() as t:
        f = Path(t) / "r.xlsx"
        wb = openpyxl.Workbook(); sh = wb.active; sh.title = SHEET
        sh.append(["ID", "Question", "Answer the bot shows", "Visual", "Needs your sign-off because", "Evidence (quote from public page)", "Source pages", "Decision", "Your edit / comment"])
        sh.append(["lp-01", "Q?", "Yes, it is.", "text", "", "", "", "Approve", None])
        sh.append(["lp-02", "Q2?", "Old.", "text", "", "", "", "Edit", "New answer here."])
        sh.append(["lp-03", "Q3?", "A.", "text", "", "", "", None, None])
        wb.save(f)
        s = read_sheet(f)
        assert [(x["id"], x["decision"], ws(x["edit"])) for x in s] == [("lp-01", "Approve", ""), ("lp-02", "Edit", "New answer here."), ("lp-03", "", "")], s
    print("import_qa_review self-check: ok")


if __name__ == "__main__":
    main()
