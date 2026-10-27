"""Apply Siddharth's decisions from the 28 Oct "not too technical" review to APPROVED website-bot pairs (sam_web_qa).

  python prototype/import_simplify_review.py --dry-run      # GETs only: read the sheet and the bank, print the plan
  python prototype/import_simplify_review.py               # apply: the database, then corpus/web/qa-citrix-lp.json
  python prototype/import_simplify_review.py FILE.xlsx --by siddharth
  python prototype/import_simplify_review.py --self-check   # offline asserts, no network

Default file: ~/Downloads/SAM-Website-Bot-Simplify-Review-2026-10-28.xlsx, sheet "Simplify review" (columns found by
header: ID, Current answer, Proposed simpler answer, Decision, Your edit). The proposals themselves (text and the exact
new visual) are corpus/web/simplify-2026-10-28.json, written with the sheet.

  Decision          What is written (the pair stays approved; approved_by, approved_at and review_by are stamped)
  Accept proposed   answer := the proposed text, and the proposal's visual change (its new block, validated)
  Edit              answer := "Your edit"; the visual is left as it is (a proposed visual change is reported, not applied)
  Keep current      nothing (a blank Decision too)

Safety, in order:
  - Every row is checked before the first write; one bad row stops the whole run (nothing is half-applied).
  - The pair must still be approved, and its answer (and its visual, when that changes) still the one the sheet showed.
  - The proposed text in the sheet must be the proposal's: to change it, put your wording in "Your edit" and pick Edit.
  - The new row must pass load_web_knowledge.py's own validator: <= 80 words and 600 characters, the emphasis rules,
    the block contract (and a stat's figure on its source page).
  - Each PATCH is filtered on status=approved. The sam_web_qa trigger bumps knowledge_version, so cached answers drop.
  - Then the same answer (and visual) goes into corpus/web/qa-citrix-lp.json, so the next loader run, the intent index
    and the review data agree with the database.
Credentials: web/.env.local (SUPABASE_URL / SUPABASE_SERVICE_KEY). Keys are never printed.
"""
from __future__ import annotations
import argparse, copy, json, sys, urllib.parse
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
import load_web_knowledge as L                                            # noqa: E402  validators + mapping
from import_qa_review import ENV, ID_RE, Db, default_by, env, review_by, ws  # noqa: E402  same stamp + credentials

REPO = HERE.parent
DEFAULT_XLSX = Path.home() / "Downloads" / "SAM-Website-Bot-Simplify-Review-2026-10-28.xlsx"
PROPOSALS = REPO / "corpus" / "web" / "simplify-2026-10-28.json"
BANK = REPO / "corpus" / "web" / "qa-citrix-lp.json"
SHEET = "Simplify review"
DECISIONS = {"accept proposed": "accept", "accept": "accept", "edit": "edit", "keep current": "keep", "keep": "keep"}
COLS = {"id": "id", "current answer": "current", "proposed simpler answer": "proposed", "decision": "decision", "your edit": "edit"}
DB_COLS = "id,status,question,aliases,answer,visual,visual_data,sources,parent_id,cta,cta_url"


def read_sheet(path: Path) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sh = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    rows = list(sh.iter_rows(values_only=True))
    wb.close()                                   # read-only mode holds the file open (Windows locks it)
    if not rows:
        return []
    head = [ws(str(h or "")).lower() for h in rows[0]]
    idx = {key: i for i, h in enumerate(head) for prefix, key in COLS.items() if h.startswith(prefix)}
    missing = [k for k in COLS.values() if k not in idx]
    if missing:
        sys.exit(f"{path.name}: no column for {', '.join(missing)} (headers: {', '.join(head)})")
    out = []
    for n, r in enumerate(rows[1:], start=2):
        get = lambda k: str(r[idx[k]]) if idx[k] < len(r) and r[idx[k]] is not None else ""
        if ws(get("id")):
            out.append({"row": n, **{k: (ws(get(k)) if k in ("id", "decision") else get(k)) for k in COLS.values()}})
    return out


def as_block(pair_id: str, question: str, bank_visual) -> dict:
    """A bank-shape visual ({type, data} or None) -> {visual, visual_data[, _stat_source]} as the loader maps it."""
    m = L.map_qa({"id": pair_id, "question": question, "visual": bank_visual}, [])
    return {k: m[k] for k in ("visual", "visual_data", "_stat_source") if k in m}


def plan(sheet: list[dict], props: dict[str, dict], db: dict[str, dict], ids: set[str], pages: dict[str, dict],
         by: str, now: datetime) -> tuple[list[dict], list[str], list[str]]:
    """-> (writes, problems, notes). A write is {id, patch, corpus, what}. Any problem means nothing is written."""
    writes, problems, notes, seen = [], [], [], set()
    stamp = {"status": "approved", "approved_by": by, "approved_at": now.astimezone(timezone.utc).isoformat(), "review_by": review_by(now)}
    for s in sheet:
        where = f"row {s['row']} ({s['id']})"
        if not s["decision"]:
            continue
        d = DECISIONS.get(s["decision"].lower())
        if not d:
            problems.append(f"{where}: unknown decision '{s['decision']}' (Accept proposed, Keep current or Edit)")
            continue
        if s["id"] in seen:
            problems.append(f"{where}: this id has a decision on another row too")
            continue
        seen.add(s["id"])
        if d == "keep":
            notes.append(f"{s['id']}: keep current, nothing written")
            continue
        p, cur = props.get(s["id"]), db.get(s["id"])
        if not ID_RE.match(s["id"]) or p is None:
            problems.append(f"{where}: no proposal for this id in {PROPOSALS.name}")
            continue
        if cur is None or cur.get("status") != "approved":
            problems.append(f"{where}: the pair is {'not in sam_web_qa' if cur is None else cur.get('status')}, not approved; review it in the Website tab")
            continue
        if ws(cur["answer"]) != ws(s["current"]) or ws(cur["answer"]) != ws(p["current_answer"]):
            problems.append(f"{where}: the approved answer changed after the sheet was made; read the current one in the Website tab")
            continue
        change_visual = d == "accept" and "proposed_visual" in p
        if change_visual:
            was = as_block(s["id"], cur["question"], p["current_visual"])
            if (cur.get("visual"), cur.get("visual_data")) != (was["visual"], was["visual_data"]):
                problems.append(f"{where}: the visual changed after the sheet was made, so the proposed visual change may not fit it")
                continue
        if d == "accept":
            if ws(s["proposed"]) != ws(p["proposed_answer"]):
                problems.append(f"{where}: the proposed text was changed in the sheet; put your wording in 'Your edit' and choose Edit")
                continue
            new = ws(p["proposed_answer"])
        else:
            new = ws(s["edit"])
            if not new:
                problems.append(f"{where}: Edit needs your wording in 'Your edit'")
                continue
            if "proposed_visual" in p:
                notes.append(f"{s['id']}: Edit changes the text only; the visual change ({p['visual_change']}) is not applied")
        row = {k: cur.get(k) for k in L.QA_COLS if k in cur}
        row["answer"] = new
        if change_visual:
            row.update(as_block(s["id"], cur["question"], p["proposed_visual"]))
        errs = L.validate_qa(row, ids, pages)
        if errs:
            problems.append(f"{where}: {'; '.join(errs)}")
            continue
        patch = {"answer": new, **({"visual": row["visual"], "visual_data": row["visual_data"]} if change_visual else {}), **stamp}
        if new == ws(cur["answer"]) and not change_visual:
            notes.append(f"{s['id']}: the new text equals the approved one, skipped")
            continue
        corpus = {"short_answer": new, **({"visual": copy.deepcopy(p["proposed_visual"])} if change_visual else {})}
        what = ("accept" if d == "accept" else "edit") + f" ({len(new.split())} words)" + (", visual changed" if change_visual else "")
        writes.append({"id": s["id"], "patch": patch, "corpus": corpus, "what": what})
    return writes, problems, notes


def write_corpus(path: Path, done: list[dict]) -> None:
    """The applied answers (and visuals) into the bank file, in its own format (indent 1, UTF-8, trailing newline)."""
    bank = json.loads(path.read_text(encoding="utf-8"))
    by_id = {q["id"]: q for q in bank}
    for w in done:
        by_id[w["id"]].update(copy.deepcopy(w["corpus"]))
    path.write_text(json.dumps(bank, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")


def corpus_refs() -> tuple[set[str], dict[str, dict]]:
    """All bank ids (parent check) and the page cards by URL (a stat's figure must be on its page)."""
    pages_raw, qa_raw, _ = L.read_corpus(None)
    return {q["id"] for q in qa_raw if q.get("id")}, {p["url"]: p for p in pages_raw if p.get("url")}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--by", help="approver login recorded in approved_by (default: the first admin in SAM_USERS)")
    ap.add_argument("--env", type=Path, default=ENV)
    ap.add_argument("--self-check", action="store_true")
    a = ap.parse_args()
    if a.self_check:
        return self_check()
    e = env(a.env)
    by = (a.by or default_by(e) or "").strip().lower()
    if not by:
        sys.exit("no approver: pass --by <your admin login>")
    if not a.xlsx.exists():
        sys.exit(f"no file {a.xlsx}")
    url, key = e.get("SUPABASE_URL"), e.get("SUPABASE_SERVICE_KEY")
    if not url or not key:
        sys.exit(f"SUPABASE_URL / SUPABASE_SERVICE_KEY missing from {a.env}")
    props = {p["id"]: p for p in json.loads(PROPOSALS.read_text(encoding="utf-8"))["pairs"]}
    sheet = read_sheet(a.xlsx)
    decided = [s for s in sheet if s["decision"]]
    print(f"{a.xlsx.name}: {len(sheet)} rows, {len(decided)} with a decision, {len(sheet) - len(decided)} undecided (left alone). Approver: {by}")
    db = Db(url, key)
    want = sorted({s["id"] for s in decided if ID_RE.match(s["id"])})
    rows = db.call("GET", f"sam_web_qa?select={DB_COLS}&id=in." + urllib.parse.quote(f"({','.join(want)})")) if want else []
    ids, pages = corpus_refs()
    writes, problems, notes = plan(sheet, props, {r["id"]: r for r in rows}, ids, pages, by, datetime.now(timezone.utc))
    for w in writes:
        print(f"  {w['id']:<8} {w['what']}")
    for n in notes:
        print(f"  note: {n}")
    if problems:
        print(f"\n{len(problems)} problem(s); nothing {'would be' if a.dry_run else 'was'} written:")
        for p in problems:
            print(f"  - {p}")
        sys.exit(1)
    if a.dry_run:
        print(f"\ndry run: {len(writes)} write(s) planned, none made. Run again without --dry-run to apply.")
        return
    kv = lambda: (db.call("GET", "sam_web_settings?select=value&key=eq.knowledge_version") or [{}])[0].get("value")
    before, done = kv(), []
    for w in writes:
        got = db.call("PATCH", f"sam_web_qa?id=eq.{urllib.parse.quote(w['id'])}&status=eq.approved", w["patch"], "return=representation")
        if len(got or []) == 1:
            done.append(w)
        else:
            print(f"  {w['id']}: no longer approved (changed in the Website tab meanwhile), skipped")
    if done:
        write_corpus(BANK, done)
    after = kv()
    print(f"applied {len(done)} of {len(writes)}; {BANK.relative_to(REPO)} updated to match. knowledge_version "
          f"{'bumped, so cached answers drop' if after != before else 'NOT bumped: check the sam_web_qa_knowledge trigger'}.\n"
          "Next: python prototype/build_intent_index.py (typed questions match the new wording), then commit the bank file.")


def self_check() -> None:
    now = datetime(2026, 10, 28, 6, 0, tzinfo=timezone.utc)
    LP = "https://accops-citrix-lp.vercel.app/"
    table = lambda cell: {"type": "compare_table", "data": {"title": "T", "columns": ["a", "b"], "rows": [{"label": "Ports", "cells": [cell]}]}}
    props = {"lp-01": {"id": "lp-01", "current_answer": "Old **one**, v7.3 and later.", "proposed_answer": "New **plain** answer for a CIO.",
                       "visual_change": "the Ports row in plain words", "current_visual": table("443"), "proposed_visual": table("An engineer confirms")},
             "lp-02": {"id": "lp-02", "current_answer": "Two.", "proposed_answer": "Two, plainly.", "visual_change": "none"},
             "lp-03": {"id": "lp-03", "current_answer": "Three.", "proposed_answer": "Three, plainly.", "visual_change": "drop the table",
                       "current_visual": table("x"), "proposed_visual": None}}
    db = {i: {"id": i, "status": "approved", "question": f"Q {i}?", "aliases": [], "answer": p["current_answer"], "sources": [LP],
              "parent_id": None, "cta": None, "cta_url": None, **as_block(i, f"Q {i}?", p.get("current_visual"))} for i, p in props.items()}
    ids = set(props)
    row = lambda i, d, edit="", n=2, cur=None, prop=None: {"row": n, "id": i, "decision": d, "edit": edit,
                                                          "current": props[i]["current_answer"] if cur is None else cur,
                                                          "proposed": props[i]["proposed_answer"] if prop is None else prop}
    w, p, n = plan([row("lp-01", "Accept proposed"), row("lp-02", "Edit", "Two, as **I** want it."), row("lp-03", "Keep current"), row("lp-03", "", n=5)],
                   props, db, ids, {}, "siddharth", now)
    assert not p, p
    assert [x["id"] for x in w] == ["lp-01", "lp-02"] and n == ["lp-03: keep current, nothing written"], (w, n)
    a, e = w[0]["patch"], w[1]["patch"]
    assert a["answer"] == "New **plain** answer for a CIO." and a["visual_data"]["rows"][0]["cells"] == ["An engineer confirms"] and a["status"] == "approved"
    assert a["approved_by"] == "siddharth" and a["approved_at"] == now.isoformat() and a["review_by"] == "2027-10-28"
    assert e["answer"] == "Two, as **I** want it." and "visual" not in e and w[1]["corpus"] == {"short_answer": "Two, as **I** want it."}
    assert w[0]["corpus"]["visual"] == props["lp-01"]["proposed_visual"] and w[0]["what"] == "accept (6 words), visual changed"
    w, p, _ = plan([row("lp-03", "accept")], props, db, ids, {}, "s", now)        # a visual dropped entirely
    assert not p and w[0]["patch"]["visual"] is None and w[0]["patch"]["visual_data"] is None and w[0]["corpus"]["visual"] is None, (w, p)
    _, _, n = plan([row("lp-01", "Edit", "Plain words only.")], props, db, ids, {}, "s", now)
    assert any("not applied" in x for x in n), n
    stale = copy.deepcopy(db); stale["lp-01"]["visual_data"]["rows"][0]["cells"] = ["8443"]
    retired = copy.deepcopy(db); retired["lp-02"]["status"] = "retired"
    for bad, d, want in [(row("lp-01", "Maybe"), db, "unknown decision"), (row("lp-01", "Accept proposed", cur="Something else."), db, "answer changed"),
                         (row("lp-01", "Accept proposed"), stale, "visual changed"), (row("lp-02", "Accept proposed"), retired, "retired"),
                         (row("lp-01", "Accept proposed", prop="My own words."), db, "proposed text was changed"),
                         (row("lp-02", "Edit", " "), db, "needs your wording"), (row("lp-02", "Edit", "w " * 81), db, "80 words"),
                         (row("lp-02", "Edit", "**a** **b** **c** **d** **e** " + "pad " * 20), db, "5 emphasised"),
                         (row("lp-02", "Edit", "Stray ** marker here."), db, "unbalanced")]:
        _, p, _ = plan([bad], props, d, ids, {}, "s", now)
        assert p and want in p[0], (want, p)
    _, p, _ = plan([row("lp-09", "Accept proposed", cur="x", prop="y")], props, db, ids, {}, "s", now)
    assert p and "no proposal" in p[0], p
    _, p, _ = plan([row("lp-02", "Edit", "One."), row("lp-02", "Keep current", n=3)], props, db, ids, {}, "s", now)
    assert p and "another row" in p[0], p
    # one bad row stops the run: the good row's write is still planned, but main() exits before writing anything
    w, p, _ = plan([row("lp-01", "Accept proposed"), row("lp-02", "Maybe", n=3)], props, db, ids, {}, "s", now)
    assert len(w) == 1 and len(p) == 1
    import openpyxl, tempfile
    with tempfile.TemporaryDirectory() as t:
        f = Path(t) / "r.xlsx"
        wb = openpyxl.Workbook(); sh = wb.active; sh.title = SHEET
        sh.append(["ID", "Question", "Current answer", "Why it's too technical", "Proposed simpler answer", "Visual change", "Decision", "Your edit"])
        sh.append(["lp-01", "Q?", "Old.", "v7.3", "New.", "none", "Accept proposed", None])
        sh.append(["lp-02", "Q2?", "Two.", "x", "Two, plainly.", "none", None, None])
        wb.save(f)
        s = read_sheet(f)
        assert [(x["id"], x["decision"], x["proposed"]) for x in s] == [("lp-01", "Accept proposed", "New."), ("lp-02", "", "Two, plainly.")], s
        bank = Path(t) / "qa.json"
        bank.write_text(json.dumps([{"id": "lp-01", "short_answer": "Old.", "visual": table("443"), "evidence": [{"quote": "é →"}]},
                                    {"id": "lp-02", "short_answer": "Two."}], ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        write_corpus(bank, [{"id": "lp-01", "corpus": {"short_answer": "New.", "visual": None}}])
        got = json.loads(bank.read_text(encoding="utf-8"))
        assert got[0]["short_answer"] == "New." and got[0]["visual"] is None and got[0]["evidence"] == [{"quote": "é →"}] and got[1]["short_answer"] == "Two."
        assert bank.read_text(encoding="utf-8").endswith("]\n") and "é →" in bank.read_text(encoding="utf-8")
    print("import_simplify_review self-check: ok")


if __name__ == "__main__":
    main()
